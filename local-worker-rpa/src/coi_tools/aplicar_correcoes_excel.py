"""
Aplicador COMPLETO de correções do Excel revisado pelo usuário.
Lê 'Relatorio_Traducao_Bancaria_Atualizado.xlsx' e:
  1. Constrói mapa: descricao_banco -> COI correto
  2. Adiciona regras ESPECÍFICAS em regras_bancarias_coi.json (alta prioridade) para TODA descrição corrigida
  3. Atualiza mapeamento_fornecedores_coi.json (por nome_sistema)
  4. Registra avisos textuais em avisos_conciliador.json
"""

import openpyxl
import json
import re
import os

ARQUIVO_EXCEL = 'Relatorio_Traducao_Bancaria_Atualizado.xlsx'
if not os.path.exists(ARQUIVO_EXCEL) and os.path.exists(os.path.join('relatorios_gerados', ARQUIVO_EXCEL)):
    ARQUIVO_EXCEL = os.path.join('relatorios_gerados', ARQUIVO_EXCEL)


# ========================
# 1. LEITURA DO EXCEL - TODAS as abas COI
# ========================
print("=" * 60)
print("Lendo correções do Excel revisado...")
wb = openpyxl.load_workbook(ARQUIVO_EXCEL)

# desc_banco -> {coi_correto, nome_sistema, status_original}
correcoes = {}
avisos = []

for sheet_name in wb.sheetnames:
    if not sheet_name.startswith('COI '):
        continue
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=2):
        if not row[1].value:
            continue
        desc_banco = str(row[1].value).strip()
        nome_sistema = str(row[4].value).strip() if row[4].value else ''
        coi_atual = str(row[5].value).strip().zfill(6) if row[5].value else ''
        col11 = row[10].value

        if col11 is None or str(col11).strip() == '':
            continue

        col11_str = str(col11).strip()
        try:
            coi_correto = str(int(float(col11_str))).zfill(6)
            eh_aviso = False
        except (ValueError, OverflowError):
            avisos.append({'desc_banco': desc_banco, 'observacao': col11_str, 'coi_atual': coi_atual})
            continue

        # Usa somente a primeira ocorrência de cada desc_banco (evita duplicatas)
        if desc_banco not in correcoes:
            correcoes[desc_banco] = {
                'nome_sistema': nome_sistema,
                'coi_atual': coi_atual,
                'coi_correto': coi_correto,
            }

mudancas = {k: v for k, v in correcoes.items() if v['coi_correto'] != v['coi_atual']}
sem_mudanca = {k: v for k, v in correcoes.items() if v['coi_correto'] == v['coi_atual']}

print(f"  → {len(mudancas)} descrições com COI diferente (correções)")
print(f"  → {len(sem_mudanca)} descrições já corretas")
print(f"  → {len(avisos)} avisos textuais")

# ========================
# 2. ATUALIZAR regras_bancarias_coi.json
#    Estratégia: TODA descrição corrigida (independente do status original)
#    recebe uma regra com padrão LITERAL de alta prioridade no início da lista.
# ========================
print("\nAtualizando regras_bancarias_coi.json...")

with open('regras_bancarias_coi.json', 'r', encoding='utf-8') as f:
    regras_bancarias = json.load(f)

# Remover regras automáticas antigas (tipo 'auto') para evitar acúmulo
regras_bancarias = [r for r in regras_bancarias if r.get('tipo') != 'auto']

# Carregar coi_data para pegar as descrições dos COIs
with open('coi_data.json', 'r', encoding='utf-8') as f:
    coi_list = json.load(f)

def get_desc_coi(codigo):
    for c in coi_list:
        if c['codigo'] == codigo:
            return c['descricao']
    return f"COI {codigo}"

novas_regras = []
for desc_banco, info in mudancas.items():
    padrao_literal = re.escape(desc_banco)
    novas_regras.append({
        "padrao": padrao_literal,
        "codigo_coi": info['coi_correto'],
        "descricao_coi": get_desc_coi(info['coi_correto']),
        "tipo": "auto"
    })

# Inserir novas regras NO INÍCIO (maior prioridade que as genéricas)
regras_bancarias_final = novas_regras + regras_bancarias

with open('regras_bancarias_coi.json', 'w', encoding='utf-8') as f:
    json.dump(regras_bancarias_final, f, indent=2, ensure_ascii=False)
print(f"  → {len(novas_regras)} regras específicas recriadas no início do arquivo.")

# ========================
# 3. ATUALIZAR mapeamento_fornecedores_coi.json
# ========================
print("\nAtualizando mapeamento_fornecedores_coi.json...")
with open('mapeamento_fornecedores_coi.json', 'r', encoding='utf-8') as f:
    mapeamento = json.load(f)

map_atualizados = 0
for cod_fornecedor, dados in mapeamento.items():
    nome_sis = dados.get('nome_identificado', '')
    for desc_banco, info in mudancas.items():
        if info['nome_sistema'] == nome_sis:
            dados['codigo_coi'] = info['coi_correto']
            dados['descricao_coi'] = get_desc_coi(info['coi_correto'])
            dados['revisado'] = True
            map_atualizados += 1
            break

with open('mapeamento_fornecedores_coi.json', 'w', encoding='utf-8') as f:
    json.dump(mapeamento, f, indent=2, ensure_ascii=False)
print(f"  → {map_atualizados} fornecedores atualizados.")

# ========================
# 4. REGISTRAR AVISOS
# ========================
print("\nRegistrando avisos...")
try:
    with open('avisos_conciliador.json', 'r', encoding='utf-8') as f:
        avisos_existentes = json.load(f)
except FileNotFoundError:
    avisos_existentes = []

descs_existentes = {a.get('desc_banco', a.get('padrao_banco', '')) for a in avisos_existentes}
novos_avisos = 0
for a in avisos:
    if a['desc_banco'] not in descs_existentes:
        avisos_existentes.append(a)
        novos_avisos += 1

with open('avisos_conciliador.json', 'w', encoding='utf-8') as f:
    json.dump(avisos_existentes, f, indent=2, ensure_ascii=False)
print(f"  → {novos_avisos} novos avisos registrados.")

# ========================
# RESUMO FINAL
# ========================
print("\n" + "=" * 60)
print("RESUMO")
print("=" * 60)
print(f"  ✅ regras_bancarias_coi.json       → {len(novas_regras)} regras específicas")
print(f"  ✅ mapeamento_fornecedores_coi.json → {map_atualizados} fornecedores")
print(f"  ⚠️  avisos_conciliador.json          → {len(avisos_existentes)} avisos totais")
if avisos:
    for a in avisos:
        print(f"     🔔 '{a['desc_banco']}': {a['observacao']}")
print("=" * 60)
