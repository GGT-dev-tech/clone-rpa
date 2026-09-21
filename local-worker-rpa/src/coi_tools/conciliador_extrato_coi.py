import csv
import json
import re
import os
import difflib
import itertools
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

# --- FUNÇÕES UTILITÁRIAS ---

def parse_valor_br(valor_str):
    if not isinstance(valor_str, str):
        return float(valor_str)
    v = valor_str.strip()
    is_negative = v.startswith('-')
    # Limpa tudo que não for dígito ou vírgula/ponto
    v = re.sub(r'[^\d,.-]', '', v)
    # Se houver ponto e vírgula, assume que ponto é milhar e vírgula é decimal
    if '.' in v and ',' in v:
        v = v.replace('.', '')
        v = v.replace(',', '.')
    # Se só houver vírgula, substitui por ponto
    elif ',' in v:
        v = v.replace(',', '.')
    
    try:
        val = float(v)
        return -abs(val) if is_negative else abs(val)
    except:
        return 0.0

def find_matching_duplicatas(target_value, duplicatas, tolerance=0.05):
    """ Busca subconjuntos de duplicatas cuja soma resulte no valor alvo usando backtracking otimizado """
    target = int(round(abs(target_value) * 100))
    tol = int(round(tolerance * 100))
    
    # Filtra duplicatas que sozinhas já excedem o valor
    valid_items = []
    for d in duplicatas:
        v = int(round(abs(d['valor_num']) * 100))
        if v <= target + tol:
            valid_items.append((d, v))
            
    # Se houver muitas duplicatas, limita o tamanho para não travar (NP-Hard)
    if len(valid_items) > 25:
        # Se for muito grande, tenta achar apenas 1 correspondência exata
        for d, v in valid_items:
            if abs(v - target) <= tol:
                return [d]
        return []
        
    # Ordena decrescente para podar a árvore mais rápido
    valid_items.sort(key=lambda x: x[1], reverse=True)
    best_match = None
    
    def backtrack(index, current_sum, current_combo):
        nonlocal best_match
        if best_match is not None:
            return
        
        if abs(current_sum - target) <= tol:
            best_match = [item[0] for item in current_combo]
            return
            
        if current_sum > target + tol or index >= len(valid_items):
            return
            
        # Limite de profundidade para combinações de no máximo 6 itens para segurança extra
        if len(current_combo) >= 6:
            return
            
        # Incluir o item atual
        backtrack(index + 1, current_sum + valid_items[index][1], current_combo + [valid_items[index]])
        # Não incluir
        backtrack(index + 1, current_sum, current_combo)

    backtrack(0, 0, [])
    return best_match if best_match else []

def clean_bank_description(desc):
    """ Remove prefixos bancários comuns para melhorar o fuzzy match """
    prefixes = [
        r'^PG\.P/INTERNET\s*-\s*',
        r'^CREDITO PIX\s*-\s*',
        r'^DEBITO PIX\s*-\s*',
        r'^PIX\s*-\s*',
        r'^TED\s*-\s*',
        r'^DOC\s*-\s*',
        r'^LIQ\.COB\.\s*',
        r'^VENDA CIELO\s*'
    ]
    cleaned = desc.upper()
    for p in prefixes:
        cleaned = re.sub(p, '', cleaned).strip()
    return cleaned

def match_supplier(bank_desc, mapping_dict):
    """ Encontra o fornecedor mais parecido no dicionário COI """
    clean_desc = clean_bank_description(bank_desc)
    
    best_match = None
    best_ratio = 0.0
    
    for cod, info in mapping_dict.items():
        nome_sis = info['nome_identificado'].upper()
        
        # Correspondência de Substring Exata (muito forte)
        if clean_desc in nome_sis or nome_sis in clean_desc:
            return cod, info
            
        ratio = difflib.SequenceMatcher(None, clean_desc, nome_sis).ratio()
        if ratio > best_ratio:
            best_ratio = ratio
            best_match = (cod, info)
            
    # Threshold de aceitação de similaridade
    if best_ratio > 0.55:
        return best_match
    return None, None

def match_bank_rule(bank_desc, regras_bancarias):
    """ Encontra se a descrição do banco bate com uma regra bancária conhecida """
    for regra in regras_bancarias:
        if re.search(regra['padrao'], bank_desc, re.IGNORECASE):
            return regra
    return None

def sugerir_codigo_coi(nome, regras_fornecedores):
    """ Tenta achar um COI sugerido usando as regras de palavras-chave, se não for encontrado no dicionário exato. """
    for cod_coi, info in regras_fornecedores.items():
        palavras = info.get("palavras_chave", [])
        for p in palavras:
            if not p: continue
            pattern = r'\b' + re.escape(p.upper()) + r'\b'
            if re.search(pattern, nome.upper()):
                return cod_coi, info['descricao_coi']
    return None, None

# --- ROTINA PRINCIPAL ---

def processar_extrato(extrato_file, mapeamento, regras_bancarias, regras_fornecedores, contas_pagas, avisos_map):
    print(f"\nProcessando arquivo: {extrato_file}")
    resultados = []
    meses_pt = {
        '01': 'Janeiro', '02': 'Fevereiro', '03': 'Marco',
        '04': 'Abril', '05': 'Maio', '06': 'Junho',
        '07': 'Julho', '08': 'Agosto', '09': 'Setembro',
        '10': 'Outubro', '11': 'Novembro', '12': 'Dezembro'
    }
    
    mes_ano = "Mes_Desconhecido"
    
    try:
        with open(extrato_file, 'r', encoding='latin-1', errors='replace') as f:
            reader = csv.reader(f, delimiter=';')
            for row in reader:
                if not row or len(row) < 5: continue
                # Layout Viacredi: Data ; Histórico ; Documento ; Valor ; C/D
                data = row[0].strip()
                if not re.match(r'\d{2}/\d{2}/\d{4}', data):
                    continue
                
                if mes_ano == "Mes_Desconhecido" and len(data) == 10:
                    mes = data[3:5]
                    ano = data[6:10]
                    nome_mes = meses_pt.get(mes, f"Mes{mes}")
                    mes_ano = f"{nome_mes}_{ano}"
                    
                historico = row[1].strip()
                doc = row[2].strip()
                valor_str = row[3].strip()
                tipo = row[4].strip()
                valor_num = parse_valor_br(valor_str)
                
                # Ignorar valores nulos ou quase zero
                if valor_num == 0: continue
                
                res = {
                    'data_banco': data,
                    'descricao_banco': historico,
                    'valor_banco': valor_num,
                    'tipo': tipo,
                    'nome_sistema': '',
                    'codigo_coi': '',
                    'descricao_coi': '',
                    'empresas': '',
                    'duplicatas_relacionadas': '',
                    'status': 'NÃO ENCONTRADO',
                    'duplicatas_brutas': []
                }
                
                # Verificar se essa descrição é um AVISO MANUAL registrado
                if historico in avisos_map:
                    res['nome_sistema'] = 'REQUER CLASSIFICAÇÃO MANUAL'
                    res['descricao_coi'] = avisos_map[historico]
                    res['status'] = 'AVISO MANUAL'
                    resultados.append(res)
                    continue
                
                # Tentar Regra Bancária primeiro
                regra_match = match_bank_rule(historico, regras_bancarias)
                
                if regra_match:
                    res['nome_sistema'] = 'TRANSAÇÃO BANCÁRIA'
                    res['codigo_coi'] = regra_match['codigo_coi']
                    res['descricao_coi'] = regra_match['descricao_coi']
                    res['status'] = 'REGRA BANCÁRIA'
                else:
                    # Buscar no dicionário de fornecedores
                    cod_fornecedor, info = match_supplier(historico, mapeamento)
                    
                    if info:
                        res['nome_sistema'] = info['nome_identificado']
                        res['codigo_coi'] = info['codigo_coi']
                        res['descricao_coi'] = info['descricao_coi']
                        
                        # Pegar as duplicatas deste fornecedor
                        dups_fornecedor = [cp for cp in contas_pagas if cp['fornecedor_codigo'] == cod_fornecedor]
                        
                        # Pegar empresas associadas
                        empresas_unicas = list(set([cp['empresa'] for cp in dups_fornecedor]))
                        res['empresas'] = ', '.join(empresas_unicas)
                        
                        if dups_fornecedor:
                            res['duplicatas_brutas'] = dups_fornecedor
                            # Tenta achar quais duplicatas compõem o valor (apenas se for Débito D)
                            if tipo == 'D':
                                matching_subset = find_matching_duplicatas(valor_num, dups_fornecedor)
                                if matching_subset:
                                    dups_strs = [f"{d['duplicata']} (R$ {d['valor']})" for d in matching_subset]
                                    res['duplicatas_relacionadas'] = ' + '.join(dups_strs)
                                    res['status'] = 'MATCH EXATO (SOMA)' if len(matching_subset) > 1 else 'MATCH EXATO'
                                else:
                                    # Se não achar soma exata, apenas lista todas as duplicatas abertas
                                    dups_strs = [f"{d['duplicata']} (R$ {d['valor']})" for d in dups_fornecedor]
                                    res['duplicatas_relacionadas'] = f"({len(dups_strs)} disponíveis): " + ' | '.join(dups_strs)
                                    res['status'] = 'FORNECEDOR ENCONTRADO (SEM SOMA EXATA)'
                            else:
                                res['status'] = 'FORNECEDOR ENCONTRADO (CRÉDITO)'
                    else:
                        # Fallback: Tentar sugerir o COI usando regras_fornecedores_coi.json
                        fallback_cod, fallback_desc = sugerir_codigo_coi(clean_bank_description(historico), regras_fornecedores)
                        if fallback_cod:
                            res['nome_sistema'] = 'FORNECEDOR NÃO CADASTRADO (SUGESTÃO)'
                            res['codigo_coi'] = fallback_cod
                            res['descricao_coi'] = fallback_desc
                            res['status'] = 'SUGESTÃO POR PALAVRA-CHAVE'
                
                resultados.append(res)
    except Exception as e:
        print(f"Erro ao processar extrato {extrato_file}: {e}")
        return
        
    # 3. Gerar Planilha Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Tradução Extrato -> SS+"
    
    headers = ["Data Banco", "Descrição Banco", "Valor (R$)", "Tipo", "Nome no Sistema", "Código COI", "Descrição COI", "Empresas", "Duplicatas Correspondentes", "Status"]
    ws.append(headers)
    
    # Estilizar cabeçalho
    header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    # Estilos de Status
    fill_green = PatternFill(start_color="C6EFCE", fill_type="solid")
    fill_yellow = PatternFill(start_color="FFEB9C", fill_type="solid")
    fill_red = PatternFill(start_color="FFC7CE", fill_type="solid")
    
    for r in resultados:
        row_data = [
            r['data_banco'], r['descricao_banco'], r['valor_banco'], r['tipo'],
            r['nome_sistema'], r['codigo_coi'], r['descricao_coi'], r['empresas'],
            r['duplicatas_relacionadas'], r['status']
        ]
        ws.append(row_data)
        
        current_row = ws.max_row
        fill_color = fill_red
        if 'MATCH EXATO' in r['status']:
            fill_color = fill_green
        elif 'FORNECEDOR ENCONTRADO' in r['status']:
            fill_color = fill_yellow
        elif r['status'] == 'REGRA BANCÁRIA':
            fill_color = PatternFill(start_color="BDD7EE", fill_type="solid")
        elif r['status'] == 'SUGESTÃO POR PALAVRA-CHAVE':
            fill_color = PatternFill(start_color="FCE4D6", fill_type="solid")
        elif r['status'] == 'AVISO MANUAL':
            fill_color = PatternFill(start_color="E2EFDA", fill_type="solid")
            
        for col in range(1, len(headers) + 1):
            ws.cell(row=current_row, column=col).fill = fill_color
            
        ws.cell(row=current_row, column=3).number_format = '#,##0.00'
        
    dims = {}
    for row in ws.rows:
        for cell in row:
            if cell.value:
                dims[cell.column_letter] = max((dims.get(cell.column_letter, 0), len(str(cell.value))))
    for col, value in dims.items():
        ws.column_dimensions[col].width = min(value + 2, 50)
        
    # --- ABA 2: DETALHES DE PENDÊNCIAS (AMARELOS) ---
    ws2 = wb.create_sheet(title="Detalhes - Pendentes")
    headers2 = ["Data Banco", "Descrição Banco", "Valor Banco", "Fornecedor", "Empresa (SS+)", "Nº Duplicata", "Vencimento", "Valor Duplicata", "Observação"]
    ws2.append(headers2)
    
    for col in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for r in resultados:
        if 'FORNECEDOR ENCONTRADO' in r['status'] and 'SEM SOMA EXATA' in r['status']:
            for d in r['duplicatas_brutas']:
                ws2.append([
                    r['data_banco'], r['descricao_banco'], r['valor_banco'],
                    r['nome_sistema'], d.get('empresa', ''), d.get('duplicata', ''),
                    d.get('vencimento', ''), parse_valor_br(d.get('valor', '0')), d.get('observacao', '')
                ])
                current_row2 = ws2.max_row
                ws2.cell(row=current_row2, column=3).number_format = '#,##0.00'
                ws2.cell(row=current_row2, column=8).number_format = '#,##0.00'
                
                for col in range(1, len(headers2) + 1):
                    ws2.cell(row=current_row2, column=col).fill = PatternFill(start_color="FFF2CC", fill_type="solid")
                    
    dims2 = {}
    for row in ws2.rows:
        for cell in row:
            if cell.value:
                dims2[cell.column_letter] = max((dims2.get(cell.column_letter, 0), len(str(cell.value))))
    for col, value in dims2.items():
        ws2.column_dimensions[col].width = min(value + 2, 45)
        
    # --- ABA 3+: AGRUPAMENTO POR CÓDIGO COI ---
    cois_encontrados = set([r['codigo_coi'] for r in resultados if r.get('codigo_coi')])
    
    for coi in sorted(list(cois_encontrados)):
        ws_coi = wb.create_sheet(title=f"COI {coi}")
        ws_coi.append(headers)
        
        for col in range(1, len(headers) + 1):
            cell = ws_coi.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        resultados_coi = [r for r in resultados if r.get('codigo_coi') == coi]
        
        for r in resultados_coi:
            row_data = [
                r['data_banco'], r['descricao_banco'], r['valor_banco'], r['tipo'],
                r['nome_sistema'], r['codigo_coi'], r['descricao_coi'], r['empresas'],
                r['duplicatas_relacionadas'], r['status']
            ]
            ws_coi.append(row_data)
            
            current_row = ws_coi.max_row
            fill_color = fill_red
            if 'MATCH EXATO' in r['status']:
                fill_color = fill_green
            elif 'FORNECEDOR ENCONTRADO' in r['status']:
                fill_color = fill_yellow
            elif r['status'] == 'REGRA BANCÁRIA':
                fill_color = PatternFill(start_color="BDD7EE", fill_type="solid")
            elif r['status'] == 'SUGESTÃO POR PALAVRA-CHAVE':
                fill_color = PatternFill(start_color="FCE4D6", fill_type="solid")
            elif r['status'] == 'AVISO MANUAL':
                fill_color = PatternFill(start_color="E2EFDA", fill_type="solid")
                
            for col in range(1, len(headers) + 1):
                ws_coi.cell(row=current_row, column=col).fill = fill_color
                
            ws_coi.cell(row=current_row, column=3).number_format = '#,##0.00'
            
        dims_coi = {}
        for row in ws_coi.rows:
            for cell in row:
                if cell.value:
                    dims_coi[cell.column_letter] = max((dims_coi.get(cell.column_letter, 0), len(str(cell.value))))
        for col, value in dims_coi.items():
            ws_coi.column_dimensions[col].width = min(value + 2, 50)
        
    pasta_saida = 'relatorios_gerados/relatorios_mensais' if os.path.exists('relatorios_gerados') else 'relatorios_mensais'
    os.makedirs(pasta_saida, exist_ok=True)
    excel_filename = f'{pasta_saida}/Relatorio_Conciliacao_{mes_ano}.xlsx'
    wb.save(excel_filename)
    print(f"  -> Relatório gerado: {excel_filename}")

def main():
    import glob
    import os
    print("Iniciando Motor Conciliador SS+ em LOTE (Extrato -> Sistema)...")
    
    # 1. Carregar Arquivos Base
    print("Carregando dicionários e regras...")
    with open('mapeamento_fornecedores_coi.json', 'r', encoding='utf-8') as f:
        mapeamento = json.load(f)
        
    with open('regras_bancarias_coi.json', 'r', encoding='utf-8') as f:
        regras_bancarias = json.load(f)
        
    with open('regras_fornecedores_coi.json', 'r', encoding='utf-8') as f:
        regras_fornecedores = json.load(f)
        
    with open('contas_pagas_data.json', 'r', encoding='utf-8') as f:
        contas_pagas = json.load(f)
        
    for cp in contas_pagas:
        cp['valor_num'] = parse_valor_br(cp['valor'])
        
    avisos_map = {}
    if os.path.exists('avisos_conciliador.json'):
        with open('avisos_conciliador.json', 'r', encoding='utf-8') as f:
            avisos_lista = json.load(f)
        avisos_map = {a['desc_banco']: a['observacao'] for a in avisos_lista}
        
    # Encontrar todos os CSVs na pasta extratos/
    arquivos_csv = glob.glob('extratos/*.csv')
    if not arquivos_csv:
        print("Nenhum arquivo CSV encontrado na pasta 'extratos/'.")
        return
        
    print(f"Encontrados {len(arquivos_csv)} arquivos para processamento.")
    
    for arquivo in sorted(arquivos_csv):
        processar_extrato(
            arquivo, mapeamento, regras_bancarias, 
            regras_fornecedores, contas_pagas, avisos_map
        )
        
    print("\nProcessamento em lote finalizado com sucesso!")

if __name__ == '__main__':
    main()
